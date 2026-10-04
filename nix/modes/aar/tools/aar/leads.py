#!/usr/bin/env python3
"""AA&R lead store: validate, merge, score, label and export (stdlib; openpyxl optional).

Leads are JSON Lines shards in $AAR_DIR/leads/*.jsonl (one shard per worker/lane/batch, so
parallel workers never edit the same file). Field reference: .roo/skills/aar-evidence-schema/SKILL.md

Usage:
  python3 tools/aar/leads.py validate            # schema + evidence checks, exit 1 on errors
  python3 tools/aar/leads.py build               # merge shards -> aar_pilot/tracker/*.csv (+ .xlsx)
  python3 tools/aar/leads.py top --n 10          # ranked JSON for the report
  python3 tools/aar/leads.py stats               # counts by lane/tab/label
"""
import argparse, csv, glob, hashlib, json, os, re, sys, urllib.parse
from datetime import datetime, timezone

AAR_DIR = os.environ.get("AAR_DIR", "aar_pilot")
BUDGET_CAP = 250
ASSET_TYPES = ["discord-bot", "open-source-tool", "browser-extension", "calculator-tool", "dev-utility",
               "template-workflow", "directory-data", "expired-domain", "wordpress-plugin", "marketplace-app",
               "browser-game", "community-forum", "resource-site", "other"]
BUNDLE_KEYS = ["source_code", "domain", "hosting", "app_ownership", "store_listing", "brand", "docs",
               "database", "community", "payment_account"]
BUNDLE_VALUES = {"Confirmed", "Likely", "Unknown"}
GATES = ["g1_owner_signal", "g2_budget_fit", "g3_asset_bundle", "g4_demand_proof", "g5_repair_path",
         "g6_lawful_path", "g7_simple_upside"]
SUBSCORES = {"demand": 25, "owner_exit": 20, "budget_fit": 15, "transferability": 15, "repair": 10,
             "upside": 10, "fit": 5}
CONFIDENCE = {"High", "Medium", "Low"}
COLUMNS = ["ID", "Asset name", "Type", "URL(s)", "Platform/source", "Owner public handle", "Public contact route",
           "Owner-exit signal (quote + link)", "Demand signals (numbers + link + date checked)", "Last update date",
           "Asking price (if listed)", "Estimated all-in cost", "Price tier", "Asset bundle", "Rights/legal notes",
           "Repair estimate", "Upside idea", "Score", "Confidence", "Status label", "Red flags",
           "NEXT ACTION (needs human approval)", "Date found"]
TABS = ["QUALIFIED", "DISCOVERY LEADS", "WATCHLIST", "REJECTED"]


def canon(url):
    u = urllib.parse.urlparse(url.strip())
    host = (u.hostname or "").lower().removeprefix("www.")
    path = re.sub(r"/+$", "", u.path) or "/"
    return f"{host}{path}".lower()


def lead_id(lead):
    return "AAR-" + hashlib.sha256(canon(lead["urls"][0]).encode()).hexdigest()[:8].upper()


def load_raw():
    rows, errors = [], []
    for path in sorted(glob.glob(os.path.join(AAR_DIR, "leads", "*.jsonl"))):
        for n, line in enumerate(open(path, encoding="utf-8"), 1):
            line = line.strip()
            if not line:
                continue
            try:
                d = json.loads(line); d["_src"] = f"{os.path.basename(path)}:{n}"; rows.append(d)
            except json.JSONDecodeError as e:
                errors.append(f"{os.path.basename(path)}:{n}: invalid JSON ({e.msg})")
    return rows, errors


def check(lead):
    e, w = [], []
    src = lead.get("_src", "?")
    for k in ("asset_name", "asset_type", "urls", "lane", "date_found"):
        if not lead.get(k):
            e.append(f"{src}: missing {k}")
    if lead.get("asset_type") and lead["asset_type"] not in ASSET_TYPES:
        e.append(f"{src}: asset_type '{lead['asset_type']}' not in {ASSET_TYPES}")
    for u in lead.get("urls") or []:
        if not str(u).startswith(("http://", "https://")):
            e.append(f"{src}: bad url {u}")
    sig = lead.get("owner_exit_signal") or {}
    if sig and not (sig.get("url") and sig.get("checked_at")):
        e.append(f"{src}: owner_exit_signal needs url and checked_at")
    for s in lead.get("demand_signals") or []:
        if not (s.get("metric") and s.get("url") and s.get("checked_at")) or s.get("value") in (None, ""):
            e.append(f"{src}: demand signal needs metric, value, url, checked_at: {s}")
    for k, v in (lead.get("asset_bundle") or {}).items():
        if k not in BUNDLE_KEYS or v not in BUNDLE_VALUES:
            e.append(f"{src}: asset_bundle {k}={v} invalid")
    for k, v in (lead.get("gates") or {}).items():
        if k not in GATES or v not in (True, False, None):
            e.append(f"{src}: gate {k}={v} invalid")
    for k, v in (lead.get("subscores") or {}).items():
        if k not in SUBSCORES or not isinstance(v, (int, float)) or not 0 <= v <= SUBSCORES[k]:
            e.append(f"{src}: subscore {k}={v} out of range 0..{SUBSCORES.get(k)}")
    if lead.get("confidence") and lead["confidence"] not in CONFIDENCE:
        e.append(f"{src}: confidence must be High/Medium/Low")
    for k in ("asking_price", "est_all_in_cost"):
        if lead.get(k) is not None and not isinstance(lead[k], (int, float)):
            e.append(f"{src}: {k} must be a number or null")
    if not lead.get("demand_signals") and not lead.get("auto_reject_reason"):
        w.append(f"{src}: no demand signal yet (cannot pass gate 4)")
    return e, w


def merge(rows):
    by = {}
    for r in rows:
        if not r.get("urls"):
            continue
        r["id"] = r.get("id") or lead_id(r)
        cur = by.get(r["id"])
        if not cur:
            by[r["id"]] = r; continue
        for k, v in r.items():  # later, non-empty values win; evidence lists are unioned
            if k == "lane":
                cur["lanes"] = sorted(set(cur.get("lanes") or [cur.get("lane")]) | {v})
                continue
            if k in ("demand_signals", "red_flags", "urls"):
                seen = {json.dumps(x, sort_keys=True) for x in cur.get(k) or []}
                cur[k] = (cur.get(k) or []) + [x for x in v or [] if json.dumps(x, sort_keys=True) not in seen]
            elif k in ("gates", "subscores", "asset_bundle"):
                cur[k] = {**(cur.get(k) or {}), **(v or {})}
            elif v not in (None, "", [], {}):
                cur[k] = v
    return list(by.values())


def tier(cost):
    if cost is None: return "Unknown"
    if cost <= 75: return "$0-75"
    if cost <= 175: return "$76-175"
    if cost <= BUDGET_CAP: return "$176-250"
    return "Over $250"


def classify(l):
    cost = l.get("est_all_in_cost")
    gates = dict(l.get("gates") or {})
    gates["g2_budget_fit"] = None if cost is None else cost <= BUDGET_CAP
    if not l.get("demand_signals"):
        gates["g4_demand_proof"] = False
    l["gates"] = gates
    l["score"] = round(sum((l.get("subscores") or {}).get(k, 0) for k in SUBSCORES))
    l["price_tier"] = tier(cost)
    sig = l.get("owner_exit_signal") or {}
    contact_q = bool(l.get("owner_handle") and l.get("contact_route") and sig.get("url"))
    notes = []
    if l.get("auto_reject_reason"):
        l["status_label"], l["tab"] = "Rejected", "REJECTED"
        return l
    if contact_q and all(gates.get(g) is True for g in GATES):
        label = "Strong candidate"
        if l["price_tier"] == "$176-250" and (l.get("subscores") or {}).get("demand", 0) < 15:
            label = "Contact-qualified"; notes.append("$176-250 tier needs demand subscore >= 15 for Strong")
        if l.get("confidence") == "Low":
            label = "Contact-qualified"; notes.append("Low confidence cannot be Strong")
    elif contact_q:
        label = "Contact-qualified"
    else:
        label = "Discovery lead"
    l["status_label"], l["label_notes"] = label, notes
    if cost is not None and cost > BUDGET_CAP:
        l["tab"] = "WATCHLIST"
    elif label in ("Strong candidate", "Contact-qualified"):
        l["tab"] = "QUALIFIED"
    else:
        l["tab"] = "DISCOVERY LEADS"
    return l


def row(l):
    sig = l.get("owner_exit_signal") or {}
    ds = "; ".join(f"{s['metric']}={s['value']} ({s['url']}, {s['checked_at']})" for s in l.get("demand_signals") or [])
    bundle = "; ".join(f"{k}:{(l.get('asset_bundle') or {}).get(k, 'Unknown')}" for k in BUNDLE_KEYS)
    money = lambda v: "" if v is None else f"${v:,.0f}"
    return [l["id"], l.get("asset_name", ""), l.get("asset_type", ""), " ".join(l.get("urls") or []),
            l.get("source_platform", ""), l.get("owner_handle", ""), l.get("contact_route", ""),
            f"\"{sig.get('quote', '')}\" {sig.get('url', '')} ({sig.get('checked_at', '')})" if sig else "",
            ds or "Unverified", l.get("last_update_date", ""), money(l.get("asking_price")),
            money(l.get("est_all_in_cost")), l["price_tier"], bundle, l.get("rights_notes", ""),
            l.get("repair_estimate", ""), l.get("upside_idea", ""), l["score"], l.get("confidence", ""),
            l["status_label"] + (f" [{'; '.join(l['label_notes'])}]" if l.get("label_notes") else ""),
            "; ".join(l.get("red_flags") or []) + (f" REJECT: {l['auto_reject_reason']}" if l.get("auto_reject_reason") else ""),
            l.get("next_action", ""), l.get("date_found", "")]


def build():
    rows, errs = load_raw()
    leads = [classify(l) for l in merge(rows)]
    leads.sort(key=lambda l: (-l["score"], l["id"]))
    out = os.path.join(AAR_DIR, "tracker"); os.makedirs(out, exist_ok=True)
    tabs = {t: [row(l) for l in leads if l["tab"] == t] for t in TABS}
    search_rows = []
    for sl in sorted(glob.glob(os.path.join(AAR_DIR, "logs", "search_log*.jsonl"))):
        for line in open(sl, encoding="utf-8"):
            try:
                d = json.loads(line)
                search_rows.append([d.get("retrieved_at"), d.get("lane"), d.get("provider"), d.get("query"),
                                    d.get("status"), d.get("n_results"), d.get("error")])
            except json.JSONDecodeError:
                pass
    for t, data in tabs.items():
        with open(os.path.join(out, t.lower().replace(" ", "_") + ".csv"), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f); w.writerow(COLUMNS); w.writerows(data)
    slcols = ["Retrieved at", "Lane", "Provider", "Query", "Status", "Results", "Error"]
    with open(os.path.join(out, "search_log.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(slcols); w.writerows(search_rows)
    json.dump(leads, open(os.path.join(out, "leads_merged.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    xlsx = None
    try:
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        for t, data in list(tabs.items()) + [("SEARCH LOG", search_rows)]:
            ws = wb.create_sheet(t[:31]); ws.append(slcols if t == "SEARCH LOG" else COLUMNS)
            for c in ws[1]: c.font = Font(bold=True)
            for r in data: ws.append(r)
            ws.freeze_panes = "A2"
        xlsx = os.path.join(out, "aar_pilot_tracker.xlsx"); wb.save(xlsx)
    except ImportError:
        pass
    return {"leads": len(leads), "by_tab": {t: len(v) for t, v in tabs.items()}, "searches": len(search_rows),
            "json_errors": errs, "xlsx": xlsx or "openpyxl not installed; CSVs only", "dir": out}


def main():
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("validate"); sub.add_parser("build"); sub.add_parser("stats")
    t = sub.add_parser("top"); t.add_argument("--n", type=int, default=10)
    a = ap.parse_args()
    if a.cmd == "validate":
        rows, errs = load_raw(); warns = []
        for r in rows:
            e, w = check(r); errs += e; warns += w
        print(json.dumps({"records": len(rows), "errors": errs, "warnings": warns[:200]}, indent=1))
        sys.exit(1 if errs else 0)
    if a.cmd == "build":
        print(json.dumps(build(), indent=1))
    elif a.cmd in ("top", "stats"):
        rows, _ = load_raw(); leads = [classify(l) for l in merge(rows)]
        if a.cmd == "top":
            pick = sorted([l for l in leads if l["tab"] in ("QUALIFIED", "DISCOVERY LEADS")], key=lambda l: -l["score"])[:a.n]
            print(json.dumps([{k: v for k, v in l.items() if k != "_src"} for l in pick], ensure_ascii=False, indent=1))
        else:
            from collections import Counter
            print(json.dumps({"total": len(leads), "tab": Counter(l["tab"] for l in leads),
                              "label": Counter(l["status_label"] for l in leads),
                              "lane": Counter(l.get("lane") for l in leads)}, indent=1))


if __name__ == "__main__":
    main()
